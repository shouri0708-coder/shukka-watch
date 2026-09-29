# -*- coding: utf-8 -*-
"""data/ の最新Excelを読み、前回状態と比較して index.html を生成する。"""
import os, re, json, sys, datetime, html
from collections import Counter

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
STATE = os.path.join(BASE, "state.json")
HIST = os.path.join(BASE, "history.json")
WATCH = os.path.join(BASE, "watch_list.txt")
OUT = os.path.join(BASE, "docs", "index.html")

try:
    import openpyxl
except ImportError:
    print("openpyxl が必要です:  py -m pip install openpyxl")
    sys.exit(1)

# ---------- Excel 読み込み ----------
KEYS = {
    "yj":     ["yj", "ｙｊ"],
    "name":   ["販売名", "品名", "商品名", "品目名", "製品名"],
    "maker":  ["製造販売業者", "企業名", "メーカー", "会社名", "製販"],
    "ship":   ["出荷対応"],
    "vol":    ["出荷量"],
    "reason": ["理由"],
    "new":    ["更新有無"],
    "outlook":["解除見込み／供給停止の解消見込み"],
    "outlook2":["在庫消尽時期", "見込み時期"],
    "generic":["成分名", "一般名"],
    "basic":  ["基礎的", "安定確保", "確保"],
    "upd":    ["⑫の情報を更新した日"],
    "upd2":   ["⑫以外の情報を更新した日"],
    "spec":   ["規格単位", "規格"],
}

YJ_RE = re.compile(r"^\d{7}[A-Z]\d{4}$")

def norm(s):
    return re.sub(r"\s+", "", str(s or "")).lower()

def detect_header(ws, max_scan=20):
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True)):
        rows.append([norm(c) for c in row])
    best, best_i = -1, 0
    for i, r in enumerate(rows):
        joined = "|".join(r)
        hits = sum(1 for ks in KEYS.values() if any(k in joined for k in ks))
        if hits > best:
            best, best_i = hits, i
    # 2段ヘッダ対応: ヘッダ行とその次の行を結合
    hdr = rows[best_i]
    nxt = rows[best_i + 1] if best_i + 1 < len(rows) else []
    if any(YJ_RE.match(c.upper()) for c in nxt):  # 次行がデータならヘッダ1段
        nxt = []
    merged = []
    for j in range(max(len(hdr), len(nxt))):
        a = hdr[j] if j < len(hdr) else ""
        b = nxt[j] if j < len(nxt) else ""
        merged.append(a + b)
    return best_i, merged

def map_columns(headers):
    col = {}
    for key, kws in KEYS.items():
        for j, h in enumerate(headers):
            if h and any(k in h for k in kws) and j not in col.values():
                col[key] = j
                break
    return col

def to_ymd(v):
    """Excelシリアル値 / datetime / 文字列 → YYYYMMDD"""
    if v is None or v == "": return ""
    if isinstance(v, datetime.datetime): return v.strftime("%Y%m%d")
    if isinstance(v, datetime.date): return v.strftime("%Y%m%d")
    t = str(v).strip()
    if re.fullmatch(r"\d{5}(\.0+)?", t):
        d = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(float(t)))
        return d.strftime("%Y%m%d")
    m = re.search(r"(\d{4})[/年.-](\d{1,2})[/月.-](\d{1,2})", t)
    if m: return "%04d%02d%02d" % tuple(int(x) for x in m.groups())
    return ""

def fmt_ymd(d):
    return f"{d[:4]}/{d[4:6]}/{d[6:8]}" if d else ""

def classify_ship(s):
    s = str(s or "")
    if "供給停止" in s or "停止" in s: return "stop"
    if "限定" in s: return "limited"
    if "通常" in s: return "normal"
    return "other" if s.strip() else ""

def load_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = max(wb.worksheets, key=lambda w: w.max_row or 0)
    hidx, headers = detect_header(ws)
    col = map_columns(headers)
    print("シート:", ws.title, "/ ヘッダ行:", hidx + 1)
    print("列マップ:", {k: headers[v][:12] for k, v in col.items()})
    items = {}
    start = hidx + 2
    # YJ列が見つからない場合は値のパターンで探す
    yj_col = col.get("yj")
    for row in ws.iter_rows(min_row=start, values_only=True):
        if not row or all(c in (None, "") for c in row):
            continue
        raw = list(row)
        cells = [("" if c is None else str(c)).strip() for c in raw]
        if yj_col is None:
            for j, c in enumerate(cells):
                if YJ_RE.match(c):
                    yj_col = j; col["yj"] = j; break
            if yj_col is None:
                continue
        yj = cells[yj_col] if yj_col < len(cells) else ""
        if not YJ_RE.match(yj):
            continue
        def g(k):
            j = col.get(k)
            return cells[j] if j is not None and j < len(cells) else ""
        ship_raw = g("ship")
        vol_raw = g("vol")
        it = {
            "yj": yj, "name": g("name"), "maker": g("maker"),
            "generic": g("generic"),
            "ship": classify_ship(ship_raw), "ship_raw": ship_raw,
            "vol": vol_raw, "reason": g("reason"),
            "outlook": (g("outlook") + (" " + (to_ymd(raw[col["outlook2"]]) and fmt_ymd(to_ymd(raw[col["outlook2"]])) or g("outlook2")) if col.get("outlook2") is not None and g("outlook2") else "")).strip(),
            "new": bool(re.search(r"new|新", g("new"), re.I)),
            "spec": g("spec"),
            "upd": to_ymd(raw[col["upd"]]) if col.get("upd") is not None and col["upd"] < len(raw) else "",
        }
        # 同一YJが複数行（包装違い等）ある場合は悪い方を優先
        rank = {"stop": 3, "limited": 2, "other": 1, "normal": 0, "": -1}
        prev = items.get(yj)
        if prev is None or rank[it["ship"]] > rank[prev["ship"]]:
            items[yj] = it
        elif it["new"]:
            prev["new"] = True
    print("品目数:", len(items), Counter(i["ship"] for i in items.values()))
    return items

# ---------- 差分 ----------
def load_json(p, default):
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    return default

def save_json(p, obj):
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)

def diff(prev_items, cur_items, date):
    ev = []
    for yj, cur in cur_items.items():
        p = prev_items.get(yj)
        c = cur["ship"]
        if p is None:
            if prev_items and c in ("limited", "stop"):
                ev.append(("new", yj, "", c))
            continue
        b = p["ship"]
        if b == c:
            continue
        if b == "normal" and c in ("limited", "stop"): t = "new"
        elif b in ("limited", "stop") and c == "normal": t = "resolved"
        elif b == "limited" and c == "stop": t = "worse"
        elif b == "stop" and c == "limited": t = "better"
        else: t = "change"
        ev.append((t, yj, b, c))
    for yj, p in prev_items.items():
        if yj not in cur_items and p["ship"] in ("limited", "stop"):
            ev.append(("removed", yj, p["ship"], ""))
    out = []
    for t, yj, b, c in ev:
        src = cur_items.get(yj) or prev_items.get(yj)
        out.append({"date": date, "type": t, "yj": yj, "name": src["name"],
                    "maker": src["maker"], "from": b, "to": c,
                    "reason": src.get("reason", ""), "outlook": src.get("outlook", "")})
    return out

# ---------- HTML ----------
LABEL = {"normal": "通常出荷", "limited": "限定出荷", "stop": "供給停止", "other": "その他", "": "-"}
TYPE_LABEL = {"new": "NEW", "resolved": "解除", "worse": "悪化", "better": "改善", "change": "変更", "removed": "掲載終了"}

def build_html(date, items, history, watch, first_run):
    active = [i for i in items.values() if i["ship"] in ("limited", "stop")]
    cnt = Counter(i["ship"] for i in items.values())
    payload = {
        "date": date, "first_run": first_run,
        "counts": {"limited": cnt["limited"], "stop": cnt["stop"], "normal": cnt["normal"], "total": len(items)},
        "active": [{k: i[k] for k in ("yj", "name", "maker", "generic", "ship", "ship_raw", "vol", "reason", "outlook", "new", "spec", "upd")} | {"since": i.get("since", "")} for i in active],
        "history": history[-3000:],
        "watch": watch,
        "label": LABEL, "tlabel": TYPE_LABEL,
    }
    js = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    return TEMPLATE.replace("__DATA__", js)

TEMPLATE = r"""<!doctype html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>出荷調整リアルタイム状況</title>
<style>
:root{--bg:#f6f7f9;--card:#fff;--tx:#1d2330;--mut:#667085;--line:#e4e7ec;--red:#d92d20;--org:#f79009;--grn:#12b76a;--blu:#2e90fa;--wat:#fff4e5}
@media(prefers-color-scheme:dark){:root{--bg:#0f1218;--card:#181c25;--tx:#e6e9f0;--mut:#98a2b3;--line:#2a3040;--wat:#3a2a12}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);font:15px/1.6 -apple-system,"Segoe UI","Hiragino Sans","Yu Gothic UI",sans-serif}
header{padding:14px 16px 6px}h1{font-size:20px;margin:0}.sub{color:var(--mut);font-size:13px}
.kpi{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:8px;padding:8px 16px}
.kpi div{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:8px 10px}
.kpi b{font-size:22px;display:block}.kpi span{font-size:12px;color:var(--mut)}
nav{display:flex;gap:6px;padding:8px 16px;overflow-x:auto}
nav button{border:1px solid var(--line);background:var(--card);color:var(--tx);padding:6px 12px;border-radius:999px;white-space:nowrap;cursor:pointer}
nav button.on{background:var(--tx);color:var(--bg);border-color:var(--tx)}
.bar{padding:0 16px 8px;display:flex;gap:8px;flex-wrap:wrap;align-items:center}#cnt{flex:1}
input[type=search],select{flex:1;min-width:160px;padding:10px 12px;font-size:16px;border:1px solid var(--line);border-radius:8px;background:var(--card);color:var(--tx);font-size:15px}
main{padding:0 16px 40px}
.item>div:first-child{display:flex;flex-wrap:wrap;align-items:center;gap:4px 6px}
.since{margin-left:auto;font-size:13px;font-weight:600;color:var(--tx);background:var(--bg);border:1px solid var(--line);border-radius:6px;padding:2px 8px;white-space:nowrap}
.since b{color:var(--red)}.since.old{color:var(--mut);font-weight:400}
.ol{font-size:13px;color:var(--mut);margin-top:2px}.ol.red{color:var(--red);font-weight:700;font-size:14px}
.item{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:10px 12px;margin-bottom:8px}
.item.w{background:var(--wat);border-color:var(--org)}
.nm{font-weight:600}.mk{color:var(--mut);font-size:13px}
.tag{display:inline-block;font-size:12px;padding:1px 8px;border-radius:999px;color:#fff;margin-right:4px;vertical-align:middle}
.t-limited{background:var(--org)}.t-stop{background:var(--red)}.t-normal{background:var(--grn)}.t-other{background:var(--mut)}
.t-new{background:var(--red)}.t-resolved{background:var(--grn)}.t-worse{background:#b42318}.t-better{background:var(--blu)}.t-change,.t-removed{background:var(--mut)}
.rs{font-size:13px;color:var(--mut)}.d{font-size:12px;color:var(--mut)}
.empty{color:var(--mut);padding:24px;text-align:center}
.more{width:100%;padding:10px;border:1px dashed var(--line);background:none;color:var(--mut);border-radius:8px;cursor:pointer}
</style></head><body>
<header><h1>出荷調整リアルタイム状況</h1><div class="sub" id="sub"></div></header>
<div class="kpi" id="kpi"></div>
<nav id="nav"></nav>
<div class="bar"><input type="search" id="q" placeholder="🔍 検索：品名・成分名・メーカー・YJコード（全角半角どちらでも）" autocomplete="off"></div>
<div class="bar"><span id="cnt" class="sub"></span><select id="days" style="flex:0 0 auto;min-width:110px"><option value="1" id="optLatest">最新</option><option value="7" selected>7日</option><option value="30">30日</option></select></div>
<main id="main"></main>
<script>
const D=__DATA__;
const $=s=>document.querySelector(s);const esc=s=>String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const days=d=>{const a=new Date(D.date.slice(0,4),D.date.slice(4,6)-1,D.date.slice(6,8)),b=new Date(d.slice(0,4),d.slice(4,6)-1,d.slice(6,8));const n=Math.round((a-b)/864e5);return n<=0?'今日から':n<30?n+'日目':Math.floor(n/30)+'か月';};
const fmt=d=>d?d.slice(0,4)+'/'+d.slice(4,6)+'/'+d.slice(6,8):'';
const watchHit=i=>D.watch.some(w=>w&&(i.yj===w||(i.name||'').includes(w)||(i.generic||'').includes(w)));
const tabs=[['watch','自店採用'],['changes','変化'],['limited','出荷調整'],['stop','出荷停止'],['resolved','解除']];
let tab=D.watch.length?'watch':'changes';
const md=(+D.date.slice(4,6))+'/'+(+D.date.slice(6,8));
$('#optLatest').textContent=md+' 公表分のみ';
$('#sub').textContent='厚労省 医療用医薬品供給状況 '+fmt(D.date)+' 公表分'+(D.first_run?'（初回取込：変化はExcelの更新日から推定）':'');
const c=D.counts;const today=D.history.filter(h=>h.date===D.date);
$('#kpi').innerHTML=`<div><b>${c.limited}</b><span>限定出荷</span></div><div><b>${c.stop}</b><span>供給停止</span></div><div><b style="color:var(--red)">${today.filter(h=>h.type==='new'||h.type==='worse').length}</b><span>${md} 新規・悪化</span></div><div><b style="color:var(--grn)">${today.filter(h=>h.type==='resolved'||h.type==='better').length}</b><span>${md} 解除・改善</span></div>`;
function tabCounts(){const co=cutoff();return{limited:D.active.filter(i=>i.ship==='limited').length,stop:D.active.filter(i=>i.ship==='stop').length,changes:D.history.filter(h=>h.date>=co&&h.type!=='resolved'&&h.type!=='better').length,resolved:D.history.filter(h=>h.date>=co&&(h.type==='resolved'||h.type==='better')).length};}
function nav(){const tabCount=tabCounts();$('#nav').innerHTML=tabs.filter(t=>t[0]!=='watch'||D.watch.length).map(t=>`<button class="${t[0]===tab?'on':''}" data-t="${t[0]}">${t[1]}${tabCount[t[0]]!=null?' '+tabCount[t[0]]:''}</button>`).join('');}
$('#nav').onclick=e=>{const b=e.target.closest('button');if(!b)return;tab=b.dataset.t;render();};
$('#q').oninput=()=>{limit=200;render();};$('#days').onchange=render;
const nz=s=>String(s||'').normalize('NFKC').toLowerCase().replace(/[\u3041-\u3096]/g,c=>String.fromCharCode(c.charCodeAt(0)+0x60)).replace(/[\s　]/g,'');
[...D.active,...D.history].forEach(o=>o._s=nz([o.name,o.generic,o.maker,o.yj,o.reason,o.spec].join('|')));
function hit(o,q){if(!q)return true;return q.split(/\s+/).every(w=>o._s.includes(w));}
function cutoff(){const n=+$('#days').value;if(n>=99999)return'0';const d=new Date(D.date.slice(0,4),D.date.slice(4,6)-1,D.date.slice(6,8));d.setDate(d.getDate()-n+1);return d.getFullYear()+String(d.getMonth()+1).padStart(2,'0')+String(d.getDate()).padStart(2,'0');}
function outlook(o){o=String(o||'').trim();if(!o)return '';const m=o.match(/^([アイウエ])[\.．]?\s*(.*)$/);const k=m?m[1]:'',rest=(m?m[2]:o).replace(/^[-－]\s*/,'').trim();
 if(k==='ア')return `<div class="ol red">▶ 再開見込み：${esc(rest||"時期未記載")}</div>`;
 if(k==='ウ')return `<div class="ol">再開見込み：未定</div>`;
 if(k==='イ')return `<div class="ol">再開見込み：なし${rest?'（'+esc(rest)+'）':''}</div>`;
 return rest?`<div class="ol">${esc(rest)}</div>`:'';}
const isNew=i=>i.since===D.date;
function itemCard(i){const r=(i.reason||'').replace(/^[０-９0-9]+[\.．]\s*/,'').replace(/^[-－]$/,'');const v=(i.vol||'').replace(/^[A-Za-zＡ-Ｚプラス]+[\.．]\s*/,'');
 return `<div class="item ${watchHit(i)?'w':''}"><div><span class="tag t-${i.ship}">${D.label[i.ship]}</span><span class="nm">${esc(i.name)}</span> <span class="mk">${esc(i.spec||'')}</span>${i.since?`<span class="since">${fmt(i.since)}〜${isNew(i)?' <b>NEW</b>':''}</span>`:`<span class="since old">2025/5以前〜</span>`}</div><div class="mk">${esc(i.maker)}${i.generic?'／'+esc(i.generic):''}</div><div class="rs">${esc(r)}${v&&v!=='出荷量通常'?'／'+esc(v):''}</div>${outlook(i.outlook)}</div>`;}
function evCard(h){return `<div class="item ${watchHit(h)?'w':''}"><div><span class="tag t-${h.type}">${D.tlabel[h.type]}</span><span class="nm">${esc(h.name)}</span></div><div class="mk">${esc(h.maker)} ／ ${h.yj}</div><div class="rs">${D.label[h.from]||'-'} → ${D.label[h.to]||'-'}${h.reason?'／'+esc(h.reason.replace(/^[０-９0-9]+[\.．]\s*/,'')):''}</div>${h.to!=='normal'?outlook(h.outlook):''}<div class="d">${fmt(h.date)}</div></div>`;}
let limit=200;
function render(){nav();$('#days').style.display=(tab==='changes'||tab==='resolved'||tab==='watch')?'':'none';const q=$('#q').value.trim().split(/[\s　]+/).filter(Boolean).map(nz).join(' ');const co=cutoff();let list=[],card;
 if(tab==='limited'||tab==='stop'){list=D.active.filter(i=>i.ship===tab&&hit(i,q)).sort((a,b)=>(isNew(b)-isNew(a))||((b.upd||'')<(a.upd||'')?-1:(b.upd||'')>(a.upd||'')?1:0));card=itemCard;}
 else if(tab==='watch'){const act=D.active.filter(i=>watchHit(i)&&hit(i,q)).map(itemCard);const ev=D.history.filter(h=>watchHit(h)&&h.date>=co&&hit(h,q)).reverse().map(evCard);
   $('#cnt').textContent=`変化 ${ev.length}件 ／ 継続中 ${act.length}件`;
   $('#main').innerHTML=(ev.length?'<h3>変化</h3>'+ev.join(''):'')+(act.length?'<h3>継続中</h3>'+act.join(''):'')||'<div class="empty">自店採用品目に該当なし</div>';return;}
 else{list=D.history.filter(h=>h.date>=co&&hit(h,q)&&(tab==='changes'?(h.type!=='resolved'&&h.type!=='better'):(h.type==='resolved'||h.type==='better'))).reverse();card=evCard;}
 $('#cnt').textContent=`該当 ${list.length.toLocaleString()}件`+(list.length>limit?`（${limit}件表示）`:'');
 $('#main').innerHTML=(list.slice(0,limit).map(card).join('')||`<div class="empty">該当なし</div>`)+(list.length>limit?`<button class="more" onclick="limit+=300;render()">さらに表示（残り${list.length-limit}）</button>`:'');}
render();
</script></body></html>"""

# ---------- main ----------
def main():
    latest = os.path.join(DATA, "latest.txt")
    if not os.path.exists(latest):
        print("先に fetch.py を実行してください"); sys.exit(1)
    date = open(latest, encoding="utf-8").read().strip()
    xlsx = os.path.join(DATA, f"{date}.xlsx")
    cur = load_excel(xlsx)
    if not cur:
        print("品目が読めませんでした。列構成を確認してください。"); sys.exit(1)

    state = load_json(STATE, {"asof": "", "items": {}})
    history = load_json(HIST, [])
    prev = state["items"]
    first_run = not prev

    if state.get("asof") == date:
        print("同じ公表日のデータは処理済み。HTMLのみ再生成します。")
        events = []
    else:
        events = diff(prev, cur, date)
        history = [h for h in history if h["date"] != date] + events

    # since（継続開始日）: 前回から引き継ぎ。前回通常→今回調整なら公表日。初見はExcelの⑬更新日（無ければ空=2025/5以前から継続）
    for yj, it in cur.items():
        p = prev.get(yj)
        if it["ship"] in ("limited", "stop"):
            if p and p.get("ship") in ("limited", "stop"):
                it["since"] = p.get("since", "") or it.get("upd", "")
            elif p:
                it["since"] = date
            else:
                it["since"] = it.get("upd") or ("" if first_run else date)
        else:
            it["since"] = ""

    # 初回のみ: ⑬更新日が直近60日の品目から履歴を補完（限定/停止→新規、通常→解除）
    if first_run:
        lim = (datetime.datetime.strptime(date, "%Y%m%d") - datetime.timedelta(days=60)).strftime("%Y%m%d")
        for yj, it in cur.items():
            u = it.get("upd")
            if not u or u < lim: continue
            t = "new" if it["ship"] in ("limited", "stop") else ("resolved" if it["ship"] == "normal" else None)
            if not t: continue
            events.append({"date": u, "type": t, "yj": yj, "name": it["name"], "maker": it["maker"],
                           "from": "", "to": it["ship"], "reason": it.get("reason", ""), "outlook": it.get("outlook", "")})
        events.sort(key=lambda e: e["date"])
        history = events[:]

    watch = []
    if os.path.exists(WATCH):
        watch = [l.strip() for l in open(WATCH, encoding="utf-8") if l.strip() and not l.startswith("#")]

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(build_html(date, cur, history, watch, first_run))
    save_json(STATE, {"asof": date, "items": cur})
    save_json(HIST, history)
    c = Counter(e["type"] for e in events)
    print(f"完了: {OUT}")
    print(f"変化: 新規{c['new']} 解除{c['resolved']} 悪化{c['worse']} 改善{c['better']} 掲載終了{c['removed']}")

if __name__ == "__main__":
    main()
